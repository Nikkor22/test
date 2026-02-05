import os
import io
from datetime import datetime
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.base import get_session
from app.models import (
    User, Subject, Teacher, ScheduleEntry, Deadline,
    Note, Material, SubjectSummary, ReminderSettings,
    TitleTemplate, GeneratedWork, UserWorkSettings
)
from app.services.reminder_service import ReminderService
from app.services.gpt_service import GPTService
from app.services.ical_sync_service import ICalSyncService

router = APIRouter(prefix="/api", tags=["api"])
gpt_service = GPTService()


# ============= Pydantic Models =============

class TeacherResponse(BaseModel):
    id: int
    name: str
    role: str
    temperament: Optional[str] = None
    preferences: Optional[str] = None
    peculiarities: Optional[str] = None
    notes: Optional[str] = None
    subject_name: str
    subject_id: int

    class Config:
        from_attributes = True


class TeacherUpdate(BaseModel):
    name: Optional[str] = None
    role: Optional[str] = None
    temperament: Optional[str] = None
    preferences: Optional[str] = None
    peculiarities: Optional[str] = None
    notes: Optional[str] = None


class TeacherCreate(BaseModel):
    subject_id: int
    name: str
    role: str = "lecturer"
    temperament: Optional[str] = None
    preferences: Optional[str] = None
    peculiarities: Optional[str] = None
    notes: Optional[str] = None


class DeadlineResponse(BaseModel):
    id: int
    title: str
    work_type: str
    description: Optional[str] = None
    gpt_description: Optional[str] = None
    deadline_date: datetime
    is_completed: bool
    subject_name: str
    subject_id: int

    class Config:
        from_attributes = True


class DeadlineCreate(BaseModel):
    subject_id: int
    title: str
    work_type: str
    description: Optional[str] = None
    deadline_date: datetime


class DeadlineUpdate(BaseModel):
    title: Optional[str] = None
    work_type: Optional[str] = None
    description: Optional[str] = None
    deadline_date: Optional[datetime] = None
    is_completed: Optional[bool] = None


class SubjectResponse(BaseModel):
    id: int
    name: str

    class Config:
        from_attributes = True


class SubjectDetailResponse(BaseModel):
    id: int
    name: str
    teachers: List[TeacherResponse] = []
    summary: Optional[str] = None

    class Config:
        from_attributes = True


class SubjectCreate(BaseModel):
    name: str


class ScheduleEntryResponse(BaseModel):
    id: int
    subject_id: int
    subject_name: str
    day_of_week: int
    start_time: str
    end_time: str
    room: Optional[str] = None
    class_type: str
    week_type: str
    teacher_name: Optional[str] = None

    class Config:
        from_attributes = True


class ScheduleEntryCreate(BaseModel):
    subject_id: int
    day_of_week: int
    start_time: str
    end_time: str
    room: Optional[str] = None
    class_type: str = "lecture"
    week_type: str = "both"
    teacher_name: Optional[str] = None


class ScheduleEntryUpdate(BaseModel):
    subject_id: Optional[int] = None
    day_of_week: Optional[int] = None
    start_time: Optional[str] = None
    end_time: Optional[str] = None
    room: Optional[str] = None
    class_type: Optional[str] = None
    week_type: Optional[str] = None
    teacher_name: Optional[str] = None


class MaterialResponse(BaseModel):
    id: int
    subject_id: int
    subject_name: str
    file_name: str
    file_type: str
    created_at: datetime

    class Config:
        from_attributes = True


class SummaryResponse(BaseModel):
    subject_id: int
    subject_name: str
    summary_text: Optional[str] = None
    generated_at: Optional[datetime] = None

    class Config:
        from_attributes = True


class ReminderSettingsResponse(BaseModel):
    hours_before: List[int]
    is_enabled: bool

    class Config:
        from_attributes = True


class ReminderSettingsUpdate(BaseModel):
    hours_before: Optional[List[int]] = None
    is_enabled: Optional[bool] = None


# === Schedule Settings ===

class ScheduleSettingsResponse(BaseModel):
    ical_url: Optional[str] = None
    last_schedule_sync: Optional[datetime] = None

    class Config:
        from_attributes = True


class ScheduleSettingsUpdate(BaseModel):
    ical_url: Optional[str] = None


class SyncResult(BaseModel):
    success: bool
    events_parsed: Optional[int] = None
    patterns_found: Optional[int] = None
    created: Optional[int] = None
    updated: Optional[int] = None
    error: Optional[str] = None


# === Title Templates ===

class TitleTemplateResponse(BaseModel):
    id: int
    name: str
    file_name: str
    is_default: bool
    created_at: datetime

    class Config:
        from_attributes = True


class TitleTemplateCreate(BaseModel):
    name: str
    is_default: bool = False


# === Generated Works ===

class GeneratedWorkResponse(BaseModel):
    id: int
    deadline_id: int
    deadline_title: str
    subject_name: str
    work_type: str
    work_number: Optional[int] = None
    file_name: Optional[str] = None
    file_type: str
    status: str  # pending, generating, ready, confirmed, sent
    scheduled_send_at: Optional[datetime] = None
    auto_send: bool
    generated_at: Optional[datetime] = None
    confirmed_at: Optional[datetime] = None
    sent_at: Optional[datetime] = None
    deadline_date: datetime

    class Config:
        from_attributes = True


class GeneratedWorkCreate(BaseModel):
    deadline_id: int
    title_template_id: Optional[int] = None
    scheduled_send_at: Optional[datetime] = None
    auto_send: bool = False


class GeneratedWorkUpdate(BaseModel):
    scheduled_send_at: Optional[datetime] = None
    auto_send: Optional[bool] = None


# === User Work Settings ===

class UserWorkSettingsResponse(BaseModel):
    reminder_days_before: List[int]
    auto_generate: bool
    generate_days_before: int
    require_confirmation: bool
    default_send_days_before: int

    class Config:
        from_attributes = True


class UserWorkSettingsUpdate(BaseModel):
    reminder_days_before: Optional[List[int]] = None
    auto_generate: Optional[bool] = None
    generate_days_before: Optional[int] = None
    require_confirmation: Optional[bool] = None
    default_send_days_before: Optional[int] = None


# === Deadline with work_number ===

class DeadlineWithWorkResponse(BaseModel):
    id: int
    title: str
    work_type: str
    work_number: Optional[int] = None
    description: Optional[str] = None
    gpt_description: Optional[str] = None
    deadline_date: datetime
    is_completed: bool
    subject_name: str
    subject_id: int
    has_generated_work: bool = False
    generated_work_status: Optional[str] = None

    class Config:
        from_attributes = True


class DeadlineCreateWithWork(BaseModel):
    subject_id: int
    title: str
    work_type: str
    work_number: Optional[int] = None
    description: Optional[str] = None
    deadline_date: datetime


# ============= Helpers =============

async def get_user_by_telegram_id(telegram_id: int, session: AsyncSession) -> User:
    result = await session.execute(
        select(User).where(User.telegram_id == telegram_id)
    )
    user = result.scalar_one_or_none()
    if not user:
        user = User(telegram_id=telegram_id)
        session.add(user)
        await session.flush()

        reminder_settings = ReminderSettings(
            user_id=user.id,
            hours_before=[72, 24, 12]
        )
        session.add(reminder_settings)
        await session.commit()
        await session.refresh(user)
    return user


# ============= Teachers =============

@router.get("/teachers", response_model=List[TeacherResponse])
async def get_teachers(
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Teacher)
        .join(Subject)
        .options(selectinload(Teacher.subject))
        .where(Subject.user_id == user.id)
        .order_by(Teacher.subject_id, Teacher.role)
    )
    teachers = result.scalars().all()
    return [
        TeacherResponse(
            id=t.id, name=t.name, role=t.role,
            temperament=t.temperament, preferences=t.preferences,
            peculiarities=t.peculiarities, notes=t.notes,
            subject_name=t.subject.name, subject_id=t.subject_id
        )
        for t in teachers
    ]


@router.get("/teachers/{teacher_id}", response_model=TeacherResponse)
async def get_teacher(
    teacher_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Teacher).join(Subject)
        .options(selectinload(Teacher.subject))
        .where(Teacher.id == teacher_id, Subject.user_id == user.id)
    )
    teacher = result.scalar_one_or_none()
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher not found")
    return TeacherResponse(
        id=teacher.id, name=teacher.name, role=teacher.role,
        temperament=teacher.temperament, preferences=teacher.preferences,
        peculiarities=teacher.peculiarities, notes=teacher.notes,
        subject_name=teacher.subject.name, subject_id=teacher.subject_id
    )


@router.post("/teachers", response_model=TeacherResponse)
async def create_teacher(
    data: TeacherCreate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Subject).where(Subject.id == data.subject_id, Subject.user_id == user.id)
    )
    subject = result.scalar_one_or_none()
    if not subject:
        raise HTTPException(status_code=404, detail="Subject not found")

    teacher = Teacher(
        subject_id=data.subject_id, name=data.name, role=data.role,
        temperament=data.temperament, preferences=data.preferences,
        peculiarities=data.peculiarities, notes=data.notes
    )
    session.add(teacher)
    await session.commit()
    await session.refresh(teacher)
    return TeacherResponse(
        id=teacher.id, name=teacher.name, role=teacher.role,
        temperament=teacher.temperament, preferences=teacher.preferences,
        peculiarities=teacher.peculiarities, notes=teacher.notes,
        subject_name=subject.name, subject_id=teacher.subject_id
    )


@router.put("/teachers/{teacher_id}", response_model=TeacherResponse)
async def update_teacher(
    teacher_id: int,
    data: TeacherUpdate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Teacher).join(Subject)
        .options(selectinload(Teacher.subject))
        .where(Teacher.id == teacher_id, Subject.user_id == user.id)
    )
    teacher = result.scalar_one_or_none()
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher not found")

    update_data = data.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(teacher, key, value)

    await session.commit()
    await session.refresh(teacher)
    return TeacherResponse(
        id=teacher.id, name=teacher.name, role=teacher.role,
        temperament=teacher.temperament, preferences=teacher.preferences,
        peculiarities=teacher.peculiarities, notes=teacher.notes,
        subject_name=teacher.subject.name, subject_id=teacher.subject_id
    )


@router.delete("/teachers/{teacher_id}")
async def delete_teacher(
    teacher_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Teacher).join(Subject)
        .where(Teacher.id == teacher_id, Subject.user_id == user.id)
    )
    teacher = result.scalar_one_or_none()
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher not found")
    await session.delete(teacher)
    await session.commit()
    return {"status": "ok"}


# ============= Deadlines =============

@router.get("/deadlines", response_model=List[DeadlineResponse])
async def get_deadlines(
    telegram_id: int = Query(...),
    show_completed: bool = Query(False),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    query = (
        select(Deadline).join(Subject)
        .options(selectinload(Deadline.subject))
        .where(Subject.user_id == user.id)
        .order_by(Deadline.deadline_date)
    )
    if not show_completed:
        query = query.where(Deadline.is_completed == False)

    result = await session.execute(query)
    deadlines = result.scalars().all()
    return [
        DeadlineResponse(
            id=d.id, title=d.title, work_type=d.work_type,
            description=d.description, gpt_description=d.gpt_description,
            deadline_date=d.deadline_date, is_completed=d.is_completed,
            subject_name=d.subject.name, subject_id=d.subject_id
        )
        for d in deadlines
    ]


@router.post("/deadlines", response_model=DeadlineResponse)
async def create_deadline(
    data: DeadlineCreate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Subject).where(Subject.id == data.subject_id, Subject.user_id == user.id)
    )
    subject = result.scalar_one_or_none()
    if not subject:
        raise HTTPException(status_code=404, detail="Subject not found")

    deadline = Deadline(
        subject_id=data.subject_id, title=data.title,
        work_type=data.work_type, description=data.description,
        deadline_date=data.deadline_date
    )
    session.add(deadline)
    await session.flush()

    # Generate GPT description
    try:
        teachers_result = await session.execute(
            select(Teacher).where(Teacher.subject_id == subject.id)
        )
        teachers = teachers_result.scalars().all()
        teacher_info = None
        if teachers:
            t = teachers[0]
            teacher_info = {
                "name": t.name, "role": t.role,
                "temperament": t.temperament, "preferences": t.preferences,
                "peculiarities": t.peculiarities
            }

        gpt_desc = await gpt_service.generate_deadline_description(
            {"subject": subject.name, "title": data.title,
             "work_type": data.work_type, "description": data.description or "",
             "deadline_date": data.deadline_date.strftime("%d.%m.%Y %H:%M")},
            teacher_info
        )
        if gpt_desc:
            deadline.gpt_description = gpt_desc
    except Exception as e:
        print(f"Error generating GPT description: {e}")

    # Create reminders
    reminder_service = ReminderService(session)
    await reminder_service.create_reminders_for_deadline(deadline, user.id)

    await session.commit()
    await session.refresh(deadline)
    return DeadlineResponse(
        id=deadline.id, title=deadline.title, work_type=deadline.work_type,
        description=deadline.description, gpt_description=deadline.gpt_description,
        deadline_date=deadline.deadline_date, is_completed=deadline.is_completed,
        subject_name=subject.name, subject_id=deadline.subject_id
    )


@router.put("/deadlines/{deadline_id}", response_model=DeadlineResponse)
async def update_deadline(
    deadline_id: int,
    data: DeadlineUpdate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Deadline).join(Subject)
        .options(selectinload(Deadline.subject))
        .where(Deadline.id == deadline_id, Subject.user_id == user.id)
    )
    deadline = result.scalar_one_or_none()
    if not deadline:
        raise HTTPException(status_code=404, detail="Deadline not found")

    update_data = data.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(deadline, key, value)

    await session.commit()
    await session.refresh(deadline)
    return DeadlineResponse(
        id=deadline.id, title=deadline.title, work_type=deadline.work_type,
        description=deadline.description, gpt_description=deadline.gpt_description,
        deadline_date=deadline.deadline_date, is_completed=deadline.is_completed,
        subject_name=deadline.subject.name, subject_id=deadline.subject_id
    )


@router.delete("/deadlines/{deadline_id}")
async def delete_deadline(
    deadline_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Deadline).join(Subject)
        .where(Deadline.id == deadline_id, Subject.user_id == user.id)
    )
    deadline = result.scalar_one_or_none()
    if not deadline:
        raise HTTPException(status_code=404, detail="Deadline not found")
    await session.delete(deadline)
    await session.commit()
    return {"status": "ok"}


# ============= Subjects =============

@router.get("/subjects", response_model=List[SubjectResponse])
async def get_subjects(
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Subject).where(Subject.user_id == user.id)
    )
    return [SubjectResponse(id=s.id, name=s.name) for s in result.scalars().all()]


@router.get("/subjects/{subject_id}", response_model=SubjectDetailResponse)
async def get_subject_detail(
    subject_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Subject)
        .options(selectinload(Subject.teachers), selectinload(Subject.summary))
        .where(Subject.id == subject_id, Subject.user_id == user.id)
    )
    subject = result.scalar_one_or_none()
    if not subject:
        raise HTTPException(status_code=404, detail="Subject not found")

    teachers = [
        TeacherResponse(
            id=t.id, name=t.name, role=t.role,
            temperament=t.temperament, preferences=t.preferences,
            peculiarities=t.peculiarities, notes=t.notes,
            subject_name=subject.name, subject_id=subject.id
        )
        for t in subject.teachers
    ]

    return SubjectDetailResponse(
        id=subject.id, name=subject.name,
        teachers=teachers,
        summary=subject.summary.summary_text if subject.summary else None
    )


@router.post("/subjects", response_model=SubjectResponse)
async def create_subject(
    data: SubjectCreate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    subject = Subject(user_id=user.id, name=data.name)
    session.add(subject)
    await session.commit()
    await session.refresh(subject)
    return SubjectResponse(id=subject.id, name=subject.name)


@router.delete("/subjects/{subject_id}")
async def delete_subject(
    subject_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Subject).where(Subject.id == subject_id, Subject.user_id == user.id)
    )
    subject = result.scalar_one_or_none()
    if not subject:
        raise HTTPException(status_code=404, detail="Subject not found")
    await session.delete(subject)
    await session.commit()
    return {"status": "ok"}


# ============= Schedule =============

@router.get("/schedule", response_model=List[ScheduleEntryResponse])
async def get_schedule(
    telegram_id: int = Query(...),
    day_of_week: Optional[int] = Query(None),
    week_type: Optional[str] = Query(None),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    query = (
        select(ScheduleEntry).join(Subject)
        .options(selectinload(ScheduleEntry.subject))
        .where(Subject.user_id == user.id)
        .order_by(ScheduleEntry.day_of_week, ScheduleEntry.start_time)
    )

    if day_of_week is not None:
        query = query.where(ScheduleEntry.day_of_week == day_of_week)

    if week_type and week_type != "both":
        query = query.where(
            (ScheduleEntry.week_type == week_type) | (ScheduleEntry.week_type == "both")
        )

    result = await session.execute(query)
    entries = result.scalars().all()
    return [
        ScheduleEntryResponse(
            id=e.id, subject_id=e.subject_id, subject_name=e.subject.name,
            day_of_week=e.day_of_week, start_time=e.start_time, end_time=e.end_time,
            room=e.room, class_type=e.class_type, week_type=e.week_type,
            teacher_name=e.teacher_name
        )
        for e in entries
    ]


@router.post("/schedule", response_model=ScheduleEntryResponse)
async def create_schedule_entry(
    data: ScheduleEntryCreate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Subject).where(Subject.id == data.subject_id, Subject.user_id == user.id)
    )
    subject = result.scalar_one_or_none()
    if not subject:
        raise HTTPException(status_code=404, detail="Subject not found")

    entry = ScheduleEntry(
        subject_id=data.subject_id, day_of_week=data.day_of_week,
        start_time=data.start_time, end_time=data.end_time,
        room=data.room, class_type=data.class_type,
        week_type=data.week_type, teacher_name=data.teacher_name
    )
    session.add(entry)
    await session.commit()
    await session.refresh(entry)
    return ScheduleEntryResponse(
        id=entry.id, subject_id=entry.subject_id, subject_name=subject.name,
        day_of_week=entry.day_of_week, start_time=entry.start_time, end_time=entry.end_time,
        room=entry.room, class_type=entry.class_type, week_type=entry.week_type,
        teacher_name=entry.teacher_name
    )


@router.put("/schedule/{entry_id}", response_model=ScheduleEntryResponse)
async def update_schedule_entry(
    entry_id: int,
    data: ScheduleEntryUpdate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(ScheduleEntry).join(Subject)
        .options(selectinload(ScheduleEntry.subject))
        .where(ScheduleEntry.id == entry_id, Subject.user_id == user.id)
    )
    entry = result.scalar_one_or_none()
    if not entry:
        raise HTTPException(status_code=404, detail="Schedule entry not found")

    update_data = data.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(entry, key, value)

    await session.commit()
    await session.refresh(entry)
    return ScheduleEntryResponse(
        id=entry.id, subject_id=entry.subject_id, subject_name=entry.subject.name,
        day_of_week=entry.day_of_week, start_time=entry.start_time, end_time=entry.end_time,
        room=entry.room, class_type=entry.class_type, week_type=entry.week_type,
        teacher_name=entry.teacher_name
    )


@router.delete("/schedule/{entry_id}")
async def delete_schedule_entry(
    entry_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(ScheduleEntry).join(Subject)
        .where(ScheduleEntry.id == entry_id, Subject.user_id == user.id)
    )
    entry = result.scalar_one_or_none()
    if not entry:
        raise HTTPException(status_code=404, detail="Schedule entry not found")
    await session.delete(entry)
    await session.commit()
    return {"status": "ok"}


# ============= Materials =============

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)


@router.get("/materials", response_model=List[MaterialResponse])
async def get_materials(
    telegram_id: int = Query(...),
    subject_id: Optional[int] = Query(None),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    query = (
        select(Material).join(Subject)
        .options(selectinload(Material.subject))
        .where(Subject.user_id == user.id)
        .order_by(Material.created_at.desc())
    )
    if subject_id:
        query = query.where(Material.subject_id == subject_id)

    result = await session.execute(query)
    materials = result.scalars().all()
    return [
        MaterialResponse(
            id=m.id, subject_id=m.subject_id, subject_name=m.subject.name,
            file_name=m.file_name, file_type=m.file_type, created_at=m.created_at
        )
        for m in materials
    ]


@router.post("/materials/upload", response_model=MaterialResponse)
async def upload_material(
    subject_id: int = Query(...),
    telegram_id: int = Query(...),
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Subject).where(Subject.id == subject_id, Subject.user_id == user.id)
    )
    subject = result.scalar_one_or_none()
    if not subject:
        raise HTTPException(status_code=404, detail="Subject not found")

    file_ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else "unknown"

    file_path = os.path.join(UPLOAD_DIR, f"{user.id}_{subject_id}_{file.filename}")
    content = await file.read()
    with open(file_path, "wb") as f:
        f.write(content)

    parsed_text = ""
    try:
        parsed_text = await parse_file_content(file_path, file_ext, content)
    except Exception as e:
        print(f"Error parsing file: {e}")

    material = Material(
        subject_id=subject_id, file_name=file.filename,
        file_type=file_ext, file_path=file_path, parsed_text=parsed_text
    )
    session.add(material)
    await session.commit()
    await session.refresh(material)

    return MaterialResponse(
        id=material.id, subject_id=material.subject_id,
        subject_name=subject.name, file_name=material.file_name,
        file_type=material.file_type, created_at=material.created_at
    )


@router.delete("/materials/{material_id}")
async def delete_material(
    material_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Material).join(Subject)
        .where(Material.id == material_id, Subject.user_id == user.id)
    )
    material = result.scalar_one_or_none()
    if not material:
        raise HTTPException(status_code=404, detail="Material not found")

    if material.file_path and os.path.exists(material.file_path):
        os.remove(material.file_path)

    await session.delete(material)
    await session.commit()
    return {"status": "ok"}


async def parse_file_content(file_path: str, file_ext: str, content: bytes) -> str:
    """Parse text content from uploaded files."""
    try:
        if file_ext == "txt":
            return content.decode("utf-8", errors="ignore")

        elif file_ext == "pdf":
            try:
                import PyPDF2
                reader = PyPDF2.PdfReader(io.BytesIO(content))
                text = ""
                for page in reader.pages:
                    text += page.extract_text() + "\n"
                return text
            except ImportError:
                return "[PDF parsing requires PyPDF2]"

        elif file_ext in ("xlsx", "xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(content))
                text = ""
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
                        if row_text.strip(" |"):
                            text += row_text + "\n"
                return text
            except ImportError:
                return "[Excel parsing requires openpyxl]"

        elif file_ext == "docx":
            try:
                import docx
                doc = docx.Document(io.BytesIO(content))
                text = "\n".join(p.text for p in doc.paragraphs)
                return text
            except ImportError:
                return "[DOCX parsing requires python-docx]"

    except Exception as e:
        return f"[Error parsing file: {e}]"

    return ""


# ============= Summaries =============

@router.get("/subjects/{subject_id}/summary", response_model=SummaryResponse)
async def get_subject_summary(
    subject_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Subject)
        .options(selectinload(Subject.summary))
        .where(Subject.id == subject_id, Subject.user_id == user.id)
    )
    subject = result.scalar_one_or_none()
    if not subject:
        raise HTTPException(status_code=404, detail="Subject not found")

    return SummaryResponse(
        subject_id=subject.id, subject_name=subject.name,
        summary_text=subject.summary.summary_text if subject.summary else None,
        generated_at=subject.summary.generated_at if subject.summary else None
    )


@router.post("/subjects/{subject_id}/summary", response_model=SummaryResponse)
async def generate_subject_summary(
    subject_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(Subject)
        .options(selectinload(Subject.summary), selectinload(Subject.materials))
        .where(Subject.id == subject_id, Subject.user_id == user.id)
    )
    subject = result.scalar_one_or_none()
    if not subject:
        raise HTTPException(status_code=404, detail="Subject not found")

    materials_text = [m.parsed_text for m in subject.materials if m.parsed_text]

    notes_result = await session.execute(
        select(Note).where(Note.user_id == user.id)
    )
    all_notes = notes_result.scalars().all()
    notes_text = [n.raw_text for n in all_notes]

    summary_text = await gpt_service.generate_subject_summary(
        subject.name, materials_text, notes_text
    )

    if subject.summary:
        subject.summary.summary_text = summary_text
        subject.summary.generated_at = datetime.utcnow()
    else:
        summary = SubjectSummary(
            subject_id=subject.id,
            summary_text=summary_text
        )
        session.add(summary)

    await session.commit()

    return SummaryResponse(
        subject_id=subject.id, subject_name=subject.name,
        summary_text=summary_text, generated_at=datetime.utcnow()
    )


# ============= Settings =============

@router.get("/settings/reminders", response_model=ReminderSettingsResponse)
async def get_reminder_settings(
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(ReminderSettings).where(ReminderSettings.user_id == user.id)
    )
    settings = result.scalar_one_or_none()
    if not settings:
        return ReminderSettingsResponse(hours_before=[72, 24, 12], is_enabled=True)
    return ReminderSettingsResponse(
        hours_before=settings.hours_before, is_enabled=settings.is_enabled
    )


@router.put("/settings/reminders", response_model=ReminderSettingsResponse)
async def update_reminder_settings(
    data: ReminderSettingsUpdate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(ReminderSettings).where(ReminderSettings.user_id == user.id)
    )
    settings = result.scalar_one_or_none()

    if not settings:
        settings = ReminderSettings(
            user_id=user.id,
            hours_before=data.hours_before or [72, 24, 12],
            is_enabled=data.is_enabled if data.is_enabled is not None else True
        )
        session.add(settings)
    else:
        if data.hours_before is not None:
            settings.hours_before = data.hours_before
        if data.is_enabled is not None:
            settings.is_enabled = data.is_enabled

    await session.commit()
    await session.refresh(settings)
    return ReminderSettingsResponse(
        hours_before=settings.hours_before, is_enabled=settings.is_enabled
    )


# ============= Schedule Settings =============

@router.get("/settings/schedule", response_model=ScheduleSettingsResponse)
async def get_schedule_settings(
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    return ScheduleSettingsResponse(
        ical_url=user.ical_url,
        last_schedule_sync=user.last_schedule_sync
    )


@router.put("/settings/schedule", response_model=ScheduleSettingsResponse)
async def update_schedule_settings(
    data: ScheduleSettingsUpdate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    if data.ical_url is not None:
        user.ical_url = data.ical_url
    await session.commit()
    await session.refresh(user)
    return ScheduleSettingsResponse(
        ical_url=user.ical_url,
        last_schedule_sync=user.last_schedule_sync
    )


@router.post("/schedule/sync", response_model=SyncResult)
async def sync_schedule(
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    user = await get_user_by_telegram_id(telegram_id, session)
    if not user.ical_url:
        return SyncResult(success=False, error="No iCal URL configured")

    sync_service = ICalSyncService(session)
    result = await sync_service.sync_user_schedule(user)
    return SyncResult(**result)


# ============= Title Templates =============

TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "templates")
os.makedirs(TEMPLATES_DIR, exist_ok=True)


@router.get("/title-templates", response_model=List[TitleTemplateResponse])
async def get_title_templates(
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Получить все шаблоны титульных листов пользователя"""
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(TitleTemplate)
        .where(TitleTemplate.user_id == user.id)
        .order_by(TitleTemplate.is_default.desc(), TitleTemplate.created_at.desc())
    )
    templates = result.scalars().all()
    return [
        TitleTemplateResponse(
            id=t.id, name=t.name, file_name=t.file_name,
            is_default=t.is_default, created_at=t.created_at
        )
        for t in templates
    ]


@router.post("/title-templates", response_model=TitleTemplateResponse)
async def upload_title_template(
    name: str = Query(...),
    is_default: bool = Query(False),
    telegram_id: int = Query(...),
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_session)
):
    """Загрузить новый шаблон титульного листа (DOCX).
    Поддерживаемые плейсхолдеры: {{subject_name}}, {{date}}, {{work_type}}, {{work_number}}, {{student_name}}, {{group_number}}
    """
    user = await get_user_by_telegram_id(telegram_id, session)

    # Validate file type
    if not file.filename.endswith(".docx"):
        raise HTTPException(status_code=400, detail="Only DOCX files are supported for title templates")

    # If setting as default, unset other defaults
    if is_default:
        await session.execute(
            select(TitleTemplate)
            .where(TitleTemplate.user_id == user.id, TitleTemplate.is_default == True)
        )
        existing_defaults = (await session.execute(
            select(TitleTemplate).where(TitleTemplate.user_id == user.id, TitleTemplate.is_default == True)
        )).scalars().all()
        for t in existing_defaults:
            t.is_default = False

    # Save file
    file_path = os.path.join(TEMPLATES_DIR, f"{user.id}_{datetime.utcnow().timestamp()}_{file.filename}")
    content = await file.read()
    with open(file_path, "wb") as f:
        f.write(content)

    template = TitleTemplate(
        user_id=user.id,
        name=name,
        file_name=file.filename,
        file_path=file_path,
        is_default=is_default
    )
    session.add(template)
    await session.commit()
    await session.refresh(template)

    return TitleTemplateResponse(
        id=template.id, name=template.name, file_name=template.file_name,
        is_default=template.is_default, created_at=template.created_at
    )


@router.put("/title-templates/{template_id}/set-default")
async def set_default_template(
    template_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Установить шаблон как шаблон по умолчанию"""
    user = await get_user_by_telegram_id(telegram_id, session)

    # Get the template
    result = await session.execute(
        select(TitleTemplate).where(TitleTemplate.id == template_id, TitleTemplate.user_id == user.id)
    )
    template = result.scalar_one_or_none()
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")

    # Unset other defaults
    existing_defaults = (await session.execute(
        select(TitleTemplate).where(TitleTemplate.user_id == user.id, TitleTemplate.is_default == True)
    )).scalars().all()
    for t in existing_defaults:
        t.is_default = False

    template.is_default = True
    await session.commit()
    return {"status": "ok"}


@router.delete("/title-templates/{template_id}")
async def delete_title_template(
    template_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Удалить шаблон титульного листа"""
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(TitleTemplate).where(TitleTemplate.id == template_id, TitleTemplate.user_id == user.id)
    )
    template = result.scalar_one_or_none()
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")

    # Delete file
    if template.file_path and os.path.exists(template.file_path):
        os.remove(template.file_path)

    await session.delete(template)
    await session.commit()
    return {"status": "ok"}


# ============= Generated Works =============

GENERATED_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "generated")
os.makedirs(GENERATED_DIR, exist_ok=True)


@router.get("/generated-works", response_model=List[GeneratedWorkResponse])
async def get_generated_works(
    telegram_id: int = Query(...),
    status: Optional[str] = Query(None),
    session: AsyncSession = Depends(get_session)
):
    """Получить все сгенерированные работы пользователя"""
    user = await get_user_by_telegram_id(telegram_id, session)

    query = (
        select(GeneratedWork)
        .join(Deadline)
        .join(Subject)
        .options(selectinload(GeneratedWork.deadline).selectinload(Deadline.subject))
        .where(Subject.user_id == user.id)
        .order_by(Deadline.deadline_date)
    )

    if status:
        query = query.where(GeneratedWork.status == status)

    result = await session.execute(query)
    works = result.scalars().all()

    return [
        GeneratedWorkResponse(
            id=w.id,
            deadline_id=w.deadline_id,
            deadline_title=w.deadline.title,
            subject_name=w.deadline.subject.name,
            work_type=w.deadline.work_type,
            work_number=w.deadline.work_number,
            file_name=w.file_name,
            file_type=w.file_type,
            status=w.status,
            scheduled_send_at=w.scheduled_send_at,
            auto_send=w.auto_send,
            generated_at=w.generated_at,
            confirmed_at=w.confirmed_at,
            sent_at=w.sent_at,
            deadline_date=w.deadline.deadline_date
        )
        for w in works
    ]


@router.get("/generated-works/{work_id}", response_model=GeneratedWorkResponse)
async def get_generated_work(
    work_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Получить информацию о конкретной сгенерированной работе"""
    user = await get_user_by_telegram_id(telegram_id, session)

    result = await session.execute(
        select(GeneratedWork)
        .join(Deadline)
        .join(Subject)
        .options(selectinload(GeneratedWork.deadline).selectinload(Deadline.subject))
        .where(GeneratedWork.id == work_id, Subject.user_id == user.id)
    )
    work = result.scalar_one_or_none()
    if not work:
        raise HTTPException(status_code=404, detail="Generated work not found")

    return GeneratedWorkResponse(
        id=work.id,
        deadline_id=work.deadline_id,
        deadline_title=work.deadline.title,
        subject_name=work.deadline.subject.name,
        work_type=work.deadline.work_type,
        work_number=work.deadline.work_number,
        file_name=work.file_name,
        file_type=work.file_type,
        status=work.status,
        scheduled_send_at=work.scheduled_send_at,
        auto_send=work.auto_send,
        generated_at=work.generated_at,
        confirmed_at=work.confirmed_at,
        sent_at=work.sent_at,
        deadline_date=work.deadline.deadline_date
    )


@router.post("/generated-works", response_model=GeneratedWorkResponse)
async def create_generated_work(
    data: GeneratedWorkCreate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Создать запись для генерации работы (статус pending)"""
    user = await get_user_by_telegram_id(telegram_id, session)

    # Validate deadline ownership
    result = await session.execute(
        select(Deadline)
        .join(Subject)
        .options(selectinload(Deadline.subject))
        .where(Deadline.id == data.deadline_id, Subject.user_id == user.id)
    )
    deadline = result.scalar_one_or_none()
    if not deadline:
        raise HTTPException(status_code=404, detail="Deadline not found")

    # Check if work already exists
    existing = await session.execute(
        select(GeneratedWork).where(GeneratedWork.deadline_id == data.deadline_id)
    )
    if existing.scalar_one_or_none():
        raise HTTPException(status_code=400, detail="Work for this deadline already exists")

    # Validate template if provided
    if data.title_template_id:
        template_result = await session.execute(
            select(TitleTemplate).where(
                TitleTemplate.id == data.title_template_id,
                TitleTemplate.user_id == user.id
            )
        )
        if not template_result.scalar_one_or_none():
            raise HTTPException(status_code=404, detail="Title template not found")

    work = GeneratedWork(
        deadline_id=data.deadline_id,
        title_template_id=data.title_template_id,
        scheduled_send_at=data.scheduled_send_at,
        auto_send=data.auto_send,
        status="pending"
    )
    session.add(work)
    await session.commit()
    await session.refresh(work)

    return GeneratedWorkResponse(
        id=work.id,
        deadline_id=work.deadline_id,
        deadline_title=deadline.title,
        subject_name=deadline.subject.name,
        work_type=deadline.work_type,
        work_number=deadline.work_number,
        file_name=work.file_name,
        file_type=work.file_type,
        status=work.status,
        scheduled_send_at=work.scheduled_send_at,
        auto_send=work.auto_send,
        generated_at=work.generated_at,
        confirmed_at=work.confirmed_at,
        sent_at=work.sent_at,
        deadline_date=deadline.deadline_date
    )


@router.post("/generated-works/{work_id}/generate")
async def generate_work_content(
    work_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Запустить генерацию содержимого работы через AI"""
    user = await get_user_by_telegram_id(telegram_id, session)

    result = await session.execute(
        select(GeneratedWork)
        .join(Deadline)
        .join(Subject)
        .options(
            selectinload(GeneratedWork.deadline).selectinload(Deadline.subject).selectinload(Subject.materials),
            selectinload(GeneratedWork.title_template)
        )
        .where(GeneratedWork.id == work_id, Subject.user_id == user.id)
    )
    work = result.scalar_one_or_none()
    if not work:
        raise HTTPException(status_code=404, detail="Generated work not found")

    if work.status not in ("pending", "ready"):
        raise HTTPException(status_code=400, detail=f"Cannot generate work with status '{work.status}'")

    # Update status
    work.status = "generating"
    await session.commit()

    try:
        # Collect materials for AI
        materials_text = []
        for material in work.deadline.subject.materials:
            if material.parsed_text:
                materials_text.append(f"=== {material.file_name} ===\n{material.parsed_text}")

        # Generate content using GPT
        from app.services.work_generator import WorkGeneratorService
        generator = WorkGeneratorService(gpt_service)

        content = await generator.generate_work_content(
            subject_name=work.deadline.subject.name,
            work_type=work.deadline.work_type,
            work_number=work.deadline.work_number,
            title=work.deadline.title,
            description=work.deadline.description,
            materials=materials_text
        )

        # Get default template if not set
        template = work.title_template
        if not template:
            template_result = await session.execute(
                select(TitleTemplate).where(
                    TitleTemplate.user_id == user.id,
                    TitleTemplate.is_default == True
                )
            )
            template = template_result.scalar_one_or_none()

        # Generate document
        file_name, file_path = await generator.create_document(
            content=content,
            subject_name=work.deadline.subject.name,
            work_type=work.deadline.work_type,
            work_number=work.deadline.work_number,
            student_name=user.first_name or "Студент",
            group_number=user.group_number or "",
            template_path=template.file_path if template else None,
            output_dir=GENERATED_DIR,
            user_id=user.id,
            deadline_id=work.deadline_id
        )

        work.content_text = content
        work.file_name = file_name
        work.file_path = file_path
        work.status = "ready"
        work.generated_at = datetime.utcnow()

        await session.commit()
        return {"status": "ok", "file_name": file_name}

    except Exception as e:
        work.status = "pending"  # Reset status on error
        await session.commit()
        raise HTTPException(status_code=500, detail=f"Generation failed: {str(e)}")


@router.post("/generated-works/{work_id}/confirm")
async def confirm_work_sending(
    work_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Подтвердить отправку работы (после этого работа будет отправлена)"""
    user = await get_user_by_telegram_id(telegram_id, session)

    result = await session.execute(
        select(GeneratedWork)
        .join(Deadline)
        .join(Subject)
        .where(GeneratedWork.id == work_id, Subject.user_id == user.id)
    )
    work = result.scalar_one_or_none()
    if not work:
        raise HTTPException(status_code=404, detail="Generated work not found")

    if work.status != "ready":
        raise HTTPException(status_code=400, detail="Work is not ready for confirmation")

    work.status = "confirmed"
    work.confirmed_at = datetime.utcnow()
    await session.commit()

    return {"status": "ok"}


@router.put("/generated-works/{work_id}", response_model=GeneratedWorkResponse)
async def update_generated_work(
    work_id: int,
    data: GeneratedWorkUpdate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Обновить настройки сгенерированной работы"""
    user = await get_user_by_telegram_id(telegram_id, session)

    result = await session.execute(
        select(GeneratedWork)
        .join(Deadline)
        .join(Subject)
        .options(selectinload(GeneratedWork.deadline).selectinload(Deadline.subject))
        .where(GeneratedWork.id == work_id, Subject.user_id == user.id)
    )
    work = result.scalar_one_or_none()
    if not work:
        raise HTTPException(status_code=404, detail="Generated work not found")

    update_data = data.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(work, key, value)

    await session.commit()
    await session.refresh(work)

    return GeneratedWorkResponse(
        id=work.id,
        deadline_id=work.deadline_id,
        deadline_title=work.deadline.title,
        subject_name=work.deadline.subject.name,
        work_type=work.deadline.work_type,
        work_number=work.deadline.work_number,
        file_name=work.file_name,
        file_type=work.file_type,
        status=work.status,
        scheduled_send_at=work.scheduled_send_at,
        auto_send=work.auto_send,
        generated_at=work.generated_at,
        confirmed_at=work.confirmed_at,
        sent_at=work.sent_at,
        deadline_date=work.deadline.deadline_date
    )


@router.delete("/generated-works/{work_id}")
async def delete_generated_work(
    work_id: int,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Удалить сгенерированную работу"""
    user = await get_user_by_telegram_id(telegram_id, session)

    result = await session.execute(
        select(GeneratedWork)
        .join(Deadline)
        .join(Subject)
        .where(GeneratedWork.id == work_id, Subject.user_id == user.id)
    )
    work = result.scalar_one_or_none()
    if not work:
        raise HTTPException(status_code=404, detail="Generated work not found")

    # Delete file
    if work.file_path and os.path.exists(work.file_path):
        os.remove(work.file_path)

    await session.delete(work)
    await session.commit()
    return {"status": "ok"}


# ============= User Work Settings =============

@router.get("/settings/work", response_model=UserWorkSettingsResponse)
async def get_work_settings(
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Получить настройки автоматической генерации работ"""
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(UserWorkSettings).where(UserWorkSettings.user_id == user.id)
    )
    settings = result.scalar_one_or_none()

    if not settings:
        return UserWorkSettingsResponse(
            reminder_days_before=[3, 1],
            auto_generate=True,
            generate_days_before=5,
            require_confirmation=True,
            default_send_days_before=1
        )

    return UserWorkSettingsResponse(
        reminder_days_before=settings.reminder_days_before,
        auto_generate=settings.auto_generate,
        generate_days_before=settings.generate_days_before,
        require_confirmation=settings.require_confirmation,
        default_send_days_before=settings.default_send_days_before
    )


@router.put("/settings/work", response_model=UserWorkSettingsResponse)
async def update_work_settings(
    data: UserWorkSettingsUpdate,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Обновить настройки автоматической генерации работ"""
    user = await get_user_by_telegram_id(telegram_id, session)
    result = await session.execute(
        select(UserWorkSettings).where(UserWorkSettings.user_id == user.id)
    )
    settings = result.scalar_one_or_none()

    if not settings:
        settings = UserWorkSettings(
            user_id=user.id,
            reminder_days_before=data.reminder_days_before or [3, 1],
            auto_generate=data.auto_generate if data.auto_generate is not None else True,
            generate_days_before=data.generate_days_before or 5,
            require_confirmation=data.require_confirmation if data.require_confirmation is not None else True,
            default_send_days_before=data.default_send_days_before or 1
        )
        session.add(settings)
    else:
        update_data = data.model_dump(exclude_unset=True)
        for key, value in update_data.items():
            setattr(settings, key, value)

    await session.commit()
    await session.refresh(settings)

    return UserWorkSettingsResponse(
        reminder_days_before=settings.reminder_days_before,
        auto_generate=settings.auto_generate,
        generate_days_before=settings.generate_days_before,
        require_confirmation=settings.require_confirmation,
        default_send_days_before=settings.default_send_days_before
    )


# ============= Deadlines with Generated Works Info =============

@router.get("/deadlines-with-works", response_model=List[DeadlineWithWorkResponse])
async def get_deadlines_with_works(
    telegram_id: int = Query(...),
    show_completed: bool = Query(False),
    session: AsyncSession = Depends(get_session)
):
    """Получить дедлайны с информацией о сгенерированных работах"""
    user = await get_user_by_telegram_id(telegram_id, session)

    query = (
        select(Deadline)
        .join(Subject)
        .options(
            selectinload(Deadline.subject),
            selectinload(Deadline.generated_work)
        )
        .where(Subject.user_id == user.id)
        .order_by(Deadline.deadline_date)
    )

    if not show_completed:
        query = query.where(Deadline.is_completed == False)

    result = await session.execute(query)
    deadlines = result.scalars().all()

    return [
        DeadlineWithWorkResponse(
            id=d.id,
            title=d.title,
            work_type=d.work_type,
            work_number=d.work_number,
            description=d.description,
            gpt_description=d.gpt_description,
            deadline_date=d.deadline_date,
            is_completed=d.is_completed,
            subject_name=d.subject.name,
            subject_id=d.subject_id,
            has_generated_work=d.generated_work is not None,
            generated_work_status=d.generated_work.status if d.generated_work else None
        )
        for d in deadlines
    ]


@router.post("/deadlines-with-work", response_model=DeadlineWithWorkResponse)
async def create_deadline_with_work(
    data: DeadlineCreateWithWork,
    telegram_id: int = Query(...),
    session: AsyncSession = Depends(get_session)
):
    """Создать дедлайн с автоматическим созданием записи для генерации работы"""
    user = await get_user_by_telegram_id(telegram_id, session)

    # Validate subject
    result = await session.execute(
        select(Subject).where(Subject.id == data.subject_id, Subject.user_id == user.id)
    )
    subject = result.scalar_one_or_none()
    if not subject:
        raise HTTPException(status_code=404, detail="Subject not found")

    # Create deadline
    deadline = Deadline(
        subject_id=data.subject_id,
        title=data.title,
        work_type=data.work_type,
        work_number=data.work_number,
        description=data.description,
        deadline_date=data.deadline_date
    )
    session.add(deadline)
    await session.flush()

    # Generate GPT description
    try:
        teachers_result = await session.execute(
            select(Teacher).where(Teacher.subject_id == subject.id)
        )
        teachers = teachers_result.scalars().all()
        teacher_info = None
        if teachers:
            t = teachers[0]
            teacher_info = {
                "name": t.name, "role": t.role,
                "temperament": t.temperament, "preferences": t.preferences,
                "peculiarities": t.peculiarities
            }

        gpt_desc = await gpt_service.generate_deadline_description(
            {"subject": subject.name, "title": data.title,
             "work_type": data.work_type, "description": data.description or "",
             "deadline_date": data.deadline_date.strftime("%d.%m.%Y %H:%M")},
            teacher_info
        )
        if gpt_desc:
            deadline.gpt_description = gpt_desc
    except Exception as e:
        print(f"Error generating GPT description: {e}")

    # Create reminders
    reminder_service = ReminderService(session)
    await reminder_service.create_reminders_for_deadline(deadline, user.id)

    # Get user work settings
    settings_result = await session.execute(
        select(UserWorkSettings).where(UserWorkSettings.user_id == user.id)
    )
    work_settings = settings_result.scalar_one_or_none()

    # Get default template
    template_result = await session.execute(
        select(TitleTemplate).where(TitleTemplate.user_id == user.id, TitleTemplate.is_default == True)
    )
    default_template = template_result.scalar_one_or_none()

    # Create generated work entry
    from datetime import timedelta
    send_days_before = work_settings.default_send_days_before if work_settings else 1
    scheduled_send_at = data.deadline_date - timedelta(days=send_days_before)

    generated_work = GeneratedWork(
        deadline_id=deadline.id,
        title_template_id=default_template.id if default_template else None,
        scheduled_send_at=scheduled_send_at,
        auto_send=not (work_settings.require_confirmation if work_settings else True),
        status="pending"
    )
    session.add(generated_work)

    await session.commit()
    await session.refresh(deadline)

    return DeadlineWithWorkResponse(
        id=deadline.id,
        title=deadline.title,
        work_type=deadline.work_type,
        work_number=deadline.work_number,
        description=deadline.description,
        gpt_description=deadline.gpt_description,
        deadline_date=deadline.deadline_date,
        is_completed=deadline.is_completed,
        subject_name=subject.name,
        subject_id=deadline.subject_id,
        has_generated_work=True,
        generated_work_status="pending"
    )
